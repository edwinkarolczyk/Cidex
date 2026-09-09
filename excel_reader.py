from __future__ import annotations

from datetime import date, datetime
import hashlib
from pathlib import Path

from openpyxl import load_workbook


class ExcelReadError(ValueError):
    pass


ALIASES = {
    "nr_zlec": {"nr zlec", "nr zlecenia", "zlecenie", "zlecenie wew", "nr zamowienia"},
    "produkt": {"produkt", "produkt / oznaczenie", "oznaczenie", "symbol", "kod produktu"},
    "ilosc": {"ilosc", "ilość", "qty", "szt", "sztuki"},
    "data_wysylki": {"data wysylki", "data wysyłki", "termin", "wysylka", "wysyłka"},
    "proces": {"proces", "operacja", "etap"},
}


def _norm(value) -> str:
    text = str(value or "").strip().casefold().replace("_", " ")
    return " ".join(text.split())


def _header_key(value):
    norm = _norm(value)
    for key, aliases in ALIASES.items():
        if norm in aliases:
            return key
    return None


def _display_date(value) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return "" if value in (None, "") else str(value).strip()


def read_excel(path: str | Path, sheet_name: str | None = None) -> dict:
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise ExcelReadError("Cidex obsługuje pliki .xlsx.")
    if not path.is_file():
        raise ExcelReadError("Nie znaleziono wskazanego pliku Excel.")
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ExcelReadError(f"Nie można otworzyć pliku Excel: {exc}") from exc
    try:
        ws = wb[sheet_name] if sheet_name else wb.active
        iterator = ws.iter_rows(values_only=True)
        header = next(iterator, None)
        if not header:
            raise ExcelReadError("Arkusz jest pusty.")
        mapping = {}
        for idx, value in enumerate(header):
            key = _header_key(value)
            if key and key not in mapping:
                mapping[key] = idx
        missing = [key for key in ("nr_zlec", "produkt", "ilosc") if key not in mapping]
        if missing:
            raise ExcelReadError("Brakuje wymaganych kolumn: " + ", ".join(missing))
        rows = []
        for row_no, values in enumerate(iterator, start=2):
            def val(key):
                idx = mapping.get(key)
                return values[idx] if idx is not None and idx < len(values) else None
            nr, product, qty = val("nr_zlec"), val("produkt"), val("ilosc")
            if nr in (None, "") and product in (None, "") and qty in (None, ""):
                continue
            rows.append({
                "source_row": row_no,
                "nr_zlec": str(nr or "").strip(),
                "produkt_input": str(product or "").strip(),
                "ilosc": qty,
                "data_wysylki": _display_date(val("data_wysylki")),
                "proces": str(val("proces") or "").strip(),
            })
        return {
            "source_name": path.name,
            "source_path": str(path.resolve()),
            "sheet": ws.title,
            "source_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "rows": rows,
        }
    finally:
        wb.close()
